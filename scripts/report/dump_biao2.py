#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dump_biao2.py — 睇【表2（審查底稿）入面到底有咩料】，判斷邊啲補得落報告。

點解要：報告嘅調整理據（期後跟進、主要發現嗰兩章）主要來自表2，但表2 加密，
到目前為止 pipeline 同 diff 嘅源頭掃描都未真正讀過佢。冇睇過內容就講「呢批數
收唔到」係冇根據。呢個 script 先出一份【結構＋覆蓋率】摘要（細，貼得返），
確認有料先再決定要唔要全文。

用法（Windows）：
    python scripts\\report\\dump_biao2.py --pw "密碼"                 # 摘要（預設）
    python scripts\\report\\dump_biao2.py --pw "密碼" --full          # 連內容（會長）
    python scripts\\report\\dump_biao2.py --pw "密碼" --full --max 60 # 每欄最多 60 字
    python scripts\\report\\dump_biao2.py data\\表2 --entity mgm --pw "密碼"

密碼亦可以唔喺 command line 畀：env KPI_XLSX_PW 或 conf/local/credentials.yml
嘅 xlsx_password。--pw 只係一次性、唔使改 conf。

出 results\\biao2_dump.txt（UTF-8）。
"""
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
try:
    import inspect_biao2 as IB
except Exception:
    IB = None
try:
    from openpyxl.utils import get_column_letter
except ImportError:
    print("✗ pip install openpyxl"); sys.exit(1)

OUT = []
# 敘述欄（要補落報告嗰啲）同 key 欄，同 inspect_biao2 一致
NARR_HINT = ["關注事項", "調整原因", "KPMG分析", "管理層解釋", "跨司", "分析意見",
             "反饋意見", "審查意見", "調整建議", "備註", "說明", "結論"]
AMT_HINT = ["調整金額", "調減", "報告投資金額", "調整後", "金額"]
KEY_HINT = ["投資項目序號", "項目序號及名稱", "項目序號", "項目編號", "項目代碼"]


def P(s=""):
    OUT.append(str(s))


def _hdr_row(ws, scan=12):
    """揾表頭行：頭幾行入面非空格最多嗰行。"""
    best, bi = -1, 1
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), 1):
        n = sum(1 for v in row if v is not None and str(v).strip())
        if n > best:
            best, bi = n, i
    return bi


def _cls(name):
    n = re.sub(r"\s+", "", str(name or ""))
    if any(k in n for k in KEY_HINT):
        return "key"
    if any(k in n for k in NARR_HINT):
        return "敘述"
    if any(k in n for k in AMT_HINT):
        return "金額"
    return ""


def run(paths, entity, full, maxlen):
    tot_files = tot_narr_cells = tot_chars = 0
    for f in paths:
        P(f"\n{'=' * 72}\n### {f.name}")
        try:
            wb = IB.load_wb(str(f)) if IB else None
            if wb is None:
                raise RuntimeError("inspect_biao2 載入唔到")
        except Exception as e:
            P(f"  ✗ 開唔到：{e}")
            continue
        tot_files += 1
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hr = _hdr_row(ws)
            hdr = [str(v).strip() if v is not None else "" for v in rows[hr - 1]]
            body = rows[hr:]
            cls = [_cls(h) for h in hdr]
            narr_i = [i for i, c in enumerate(cls) if c == "敘述"]
            P(f"\n  ── sheet「{ws.title}」　{len(body)} 行 x {len(hdr)} 欄"
              f"（表頭喺第 {hr} 行）")
            if not any(hdr):
                P("     （冇表頭，跳過）"); continue
            P(f"     {'欄':>4} {'類':<5} {'欄名':<30} 有值行 / 總字數")
            for i, h in enumerate(hdr):
                if not h:
                    continue
                vals = [r[i] for r in body if i < len(r) and r[i] is not None
                        and str(r[i]).strip()]
                chars = sum(len(str(v).strip()) for v in vals)
                if cls[i] == "敘述":
                    tot_narr_cells += len(vals); tot_chars += chars
                mark = {"key": "key", "敘述": "★敘述", "金額": "金額"}.get(cls[i], "")
                if not vals and not mark:
                    continue
                P(f"     {get_column_letter(i + 1):>4} {mark:<5} {h[:30]:<30} "
                  f"{len(vals):>4} / {chars:,}")
            if full and narr_i:
                kc = next((i for i, c in enumerate(cls) if c == "key"), 0)
                P(f"\n     —— 內容（每欄最多 {maxlen} 字）——")
                for r in body:
                    if kc >= len(r) or r[kc] is None:
                        continue
                    segs = []
                    for i in narr_i:
                        if i < len(r) and r[i] is not None and str(r[i]).strip():
                            t = re.sub(r"\s+", " ", str(r[i]).strip())
                            segs.append(f"{hdr[i]}：{t[:maxlen]}")
                    if segs:
                        P(f"     [{str(r[kc]).strip()[:24]}] " + " ｜ ".join(segs))
        try:
            wb.close()
        except Exception:
            pass
    P(f"\n{'=' * 72}")
    P(f"→ 開到 {tot_files}/{len(paths)} 個檔；敘述欄合共 {tot_narr_cells:,} 格、"
      f"{tot_chars:,} 字")
    if tot_chars:
        P("  （報告『期後跟進』同『主要發現』兩章 golden 合共約 20,700 字 —— "
          "上面呢啲就係可以餵落去嘅原料）")


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    av = sys.argv[1:]
    if "--pw" in av:
        import os
        os.environ["KPI_XLSX_PW"] = av[av.index("--pw") + 1]
        if IB:                                   # PASSWORD 喺 import 嗰陣已經定咗
            IB.PASSWORD = av[av.index("--pw") + 1]
    ent = (av[av.index("--entity") + 1] if "--entity" in av else "mgm").lower()
    full = "--full" in av
    maxlen = int(av[av.index("--max") + 1]) if "--max" in av else 160
    pos = [a for a in av if not a.startswith("--")
           and a not in {av[av.index(k) + 1] for k in ("--pw", "--entity", "--max") if k in av}]
    root = Path(pos[0]) if pos else Path("data/表2")
    if not root.exists():
        print(f"✗ 搵唔到 {root}"); return
    paths = sorted(x for x in ([root] if root.is_file() else root.rglob("*.xls*"))
                   if not x.name.startswith("~$") and ent in x.name.lower())
    if not paths:
        print(f"✗ {root} 入面冇 match「{ent}」嘅 xlsx"); return
    run(paths, ent, full, maxlen)
    txt = "\n".join(OUT)
    dest = Path("results") if Path("results").is_dir() else Path(".")
    fo = dest / "biao2_dump.txt"
    fo.write_text(txt, encoding="utf-8")
    print(txt)
    print(f"\n✓ 已寫 {fo}（UTF-8，{len(OUT)} 行）")
    print("⚠ 內容係客戶審查底稿 —— results/ 已 gitignore，唔好 commit")


if __name__ == "__main__":
    main()
